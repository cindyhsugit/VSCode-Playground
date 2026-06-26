import csv

import os
print(os.getcwd())
import csv

# with open(r"c:/Users/cindy/Documents/VSCode Playground/src/data.csv") as file:
#     reader = csv.reader(file)
#     for row in reader:
#         print(row)

from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

ws["A1"] = "Name"
ws["B1"] = "Age"

ws["A2"] = "Alice"
ws["B2"] = 22

ws["A3"] = "Bob"
ws["B3"] = 25

wb.save("c:/Users/cindy/Documents/VSCode Playground/src/data.xlsx")

import pandas as pd

df = pd.read_excel("c:/Users/cindy/Documents/VSCode Playground/src/data.xlsx", engine="openpyxl")
#print(df)

df = pd.read_csv("c:/Users/cindy/Documents/VSCode Playground/src/data.csv")
df.dropna(inplace=True)
#print(df)
aapl = df[df["Ticker"] == "AAPL"]
print(aapl)
